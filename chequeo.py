import os, requests, smtplib, io
from datetime import date, datetime, timezone, timedelta
from email.mime.text import MIMEText
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
from google.oauth2.service_account import Credentials
import openpyxl
import json

# --- Credenciales ---
CONNECTEAM_API_KEY = os.environ["CONNECTEAM_API_KEY"]
GMAIL_SENDER       = os.environ["GMAIL_SENDER"]
GMAIL_RECIPIENT    = os.environ["GMAIL_RECIPIENT"]
GMAIL_APP_PASSWORD = os.environ["GMAIL_APP_PASSWORD"]
SA_JSON            = os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"]

REGISTRO_ID = "1duUISVaQ4Lk9djwpOGXY88_jjSVsKn6N"
BA_TZ = timezone(timedelta(hours=-3))

MESES_ES = {
    1:"Enero",2:"Febrero",3:"Marzo",4:"Abril",5:"Mayo",6:"Junio",
    7:"Julio",8:"Agosto",9:"Septiembre",10:"Octubre",11:"Noviembre",12:"Diciembre"
}
P1      = ["bilder","vonderk","esparza","correa","triunvirato 5375"]
P2      = ["amenábar 3208","ciudad de la paz","core oficina","conesa 2958","vonderk depósito"]
IGNORAR = ["walter benitez","rodrigo martinez"]

MIME_XLSX  = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DRIVE_OPTS = dict(supportsAllDrives=True, includeItemsFromAllDrives=True)

# --- Auth Google Drive ---
creds = Credentials.from_service_account_info(
    json.loads(SA_JSON),
    scopes=["https://www.googleapis.com/auth/drive"]
)
drive = build("drive", "v3", credentials=creds)

hoy            = date.today()
nombre_archivo = hoy.strftime("%d/%m/%Y") + ".xlsx"
nombre_carpeta = f"{MESES_ES[hoy.month]} {hoy.year}"

# --- Buscar carpeta del mes ---
q = f"name='{nombre_carpeta}' and mimeType='application/vnd.google-apps.folder' and '{REGISTRO_ID}' in parents and trashed=false"
files = drive.files().list(q=q, fields="files(id)", **DRIVE_OPTS).execute().get("files", [])
if not files:
    raise FileNotFoundError(f"No existe la carpeta '{nombre_carpeta}'")
carpeta_mes_id = files[0]["id"]

# --- Buscar archivo del día ---
q2 = f"name='{nombre_archivo}' and '{carpeta_mes_id}' in parents and trashed=false"
files2 = drive.files().list(q=q2, fields="files(id)", **DRIVE_OPTS).execute().get("files", [])
if not files2:
    raise FileNotFoundError(f"No existe el archivo '{nombre_archivo}'")
archivo_id = files2[0]["id"]
print(f"Archivo encontrado: {nombre_archivo}")

# --- Descargar archivo ---
buf = io.BytesIO()
dl  = MediaIoBaseDownload(buf, drive.files().get_media(fileId=archivo_id, supportsAllDrives=True))
done = False
while not done:
    _, done = dl.next_chunk()
buf.seek(0)
wb = openpyxl.load_workbook(buf)
ws = wb.active

# --- Connecteam ---
ct = {"X-API-KEY": CONNECTEAM_API_KEY, "Content-Type": "application/json"}

# Jobs con paginación
job_nombre = {}
offset = 0
while True:
    r_jobs = requests.get("https://api.connecteam.com/jobs/v1/jobs", headers=ct,
                          params={"limit": 50, "offset": offset})
    jobs_list = r_jobs.json().get("data", {}).get("jobs", [])
    for j in jobs_list:
        if "jobId" in j and "title" in j:
            job_nombre[j["jobId"]] = j["title"]
    if len(jobs_list) < 50:
        break
    offset += 50
print(f"Jobs cargados: {len(job_nombre)}")

# Usuarios
user_nombre = {}
r_users = requests.get("https://api.connecteam.com/users/v1/users", headers=ct)
for u in r_users.json().get("data", {}).get("users", []):
    if "userId" in u:
        user_nombre[u["userId"]] = f"{u.get('firstName','')} {u.get('lastName','')}".strip()
print(f"Usuarios cargados: {len(user_nombre)}")

# Scheduler y turnos del día
scheduler_id = requests.get("https://api.connecteam.com/scheduler/v1/schedulers", headers=ct).json().get("data",{}).get("schedulers",[{}])[0].get("schedulerId")
hoy_ts_start = int(datetime.combine(hoy, datetime.min.time()).timestamp())
hoy_ts_end   = int(datetime.combine(hoy, datetime.max.time()).timestamp())
turnos = requests.get(
    f"https://api.connecteam.com/scheduler/v1/schedulers/{scheduler_id}/shifts",
    headers=ct, params={"startTime": hoy_ts_start, "endTime": hoy_ts_end}
).json().get("data", {}).get("shifts", [])
print(f"Turnos encontrados: {len(turnos)}")

# Fichajes
timeclock_id = requests.get("https://api.connecteam.com/time-clock/v1/time-clocks", headers=ct).json().get("data",{}).get("timeClocks",[{}])[0].get("id")
fichajes_por_usuario = {}
for ud in requests.get(
    f"https://api.connecteam.com/time-clock/v1/time-clocks/{timeclock_id}/time-activities",
    headers=ct, params={"startDate": hoy.isoformat(), "endDate": hoy.isoformat()}
).json().get("data", {}).get("timeActivitiesByUsers", []):
    uid = ud.get("userId")
    if ud.get("shifts"):
        fichajes_por_usuario[uid] = ud["shifts"]
print(f"Usuarios con fichajes: {list(fichajes_por_usuario.keys())}")

# Debug
print("=== SERVICIOS EN PLANILLA (Col C, fila 5+) ===")
for row in ws.iter_rows(min_row=5):
    if row[2].value:
        print(f"  {repr(str(row[2].value).strip())}")

# --- Cruzar y completar ---
ausencias, tardanzas = [], []
cubiertos, total = 0, 0

for row in ws.iter_rows(min_row=5):
    servicio = row[2].value  # Col C
    if not servicio:
        continue
    if row[4].value:  # Col E ya completada
        continue

    servicio_str   = str(servicio).strip()
    servicio_lower = servicio_str.lower()

    turno = next((t for t in turnos
                  if job_nombre.get(t.get("jobId",""),"").strip().lower() == servicio_lower), None)

    if not turno:
        continue  # Sin turno hoy para este servicio

    uid         = turno.get("assignedUserIds", [None])[0]
    operario    = user_nombre.get(uid, f"ID:{uid}")
    turno_jid   = turno.get("jobId","")
    turno_start = turno.get("startTime", 0)

    if str(operario).strip().lower() in IGNORAR:
        continue

    row[3].value = operario  # Col D
    total += 1

    fichaje = next((f for f in fichajes_por_usuario.get(uid, []) if f.get("jobId") == turno_jid), None)
    if not fichaje:
        fichaje = next((f for f in fichajes_por_usuario.get(uid, [])
                        if abs(f.get("start", {}).get("timestamp", 0) - turno_start) < 7200), None)

    if not fichaje:
        row[4].value = "-"
        row[5].value = "NO"
        row[6].value = "No fichó"
        prioridad = "P1" if any(p in servicio_lower for p in P1) else \
                    "P2" if any(p in servicio_lower for p in P2) else "P3"
        row[8].value = prioridad
        ausencias.append({"nombre": operario, "servicio": servicio_str,
                          "horario": datetime.fromtimestamp(turno_start, BA_TZ).strftime("%H:%M"),
                          "prioridad": prioridad})
    else:
        clock_ts  = fichaje.get("start", {}).get("timestamp", 0)
        diff      = (clock_ts - turno_start) / 60
        hora_real = datetime.fromtimestamp(clock_ts, BA_TZ).strftime("%H:%M")
        hora_prog = datetime.fromtimestamp(turno_start, BA_TZ).strftime("%H:%M")
        if diff > 10:
            row[4].value = "X"
            row[5].value = "TARDE"
            row[6].value = hora_real
            tardanzas.append({"nombre": operario, "servicio": servicio_str,
                              "hora_prog": hora_prog, "hora_real": hora_real})
        else:
            row[4].value = "X"
            row[5].value = "OK"
            cubiertos += 1

# --- Subir planilla ---
buf2 = io.BytesIO()
wb.save(buf2)
buf2.seek(0)
drive.files().update(fileId=archivo_id,
                     media_body=MediaIoBaseUpload(buf2, mimetype=MIME_XLSX),
                     supportsAllDrives=True).execute()
print("Planilla actualizada en Drive")

# --- Mail ---
hora_actual  = datetime.now(BA_TZ).strftime("%H:%M")
fecha_actual = hoy.strftime("%d/%m/%Y")

lineas_aus = "\n".join(
    f"- {a['nombre']} | {a['servicio']} | {a['horario']} | {a['prioridad']}"
    for a in ausencias) or "Ninguna"
lineas_tar = "\n".join(
    f"- {t['nombre']} | {t['servicio']} | {t['hora_prog']} → {t['hora_real']}"
    for t in tardanzas) or "Ninguna"

cuerpo = f"""AUSENCIAS:
{lineas_aus}

TARDANZAS:
{lineas_tar}

RESUMEN:
Cubiertos: {cubiertos} / Total: {total}
P1 sin cubrir: {sum(1 for a in ausencias if a['prioridad']=='P1')}
P2 sin cubrir: {sum(1 for a in ausencias if a['prioridad']=='P2')}
"""

msg = MIMEText(cuerpo, "plain", "utf-8")
msg["Subject"] = f"Asistencia {hora_actual} — {fecha_actual}"
msg["From"]    = GMAIL_SENDER
msg["To"]      = GMAIL_RECIPIENT

with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
    server.login(GMAIL_SENDER, GMAIL_APP_PASSWORD)
    server.send_message(msg)

print(f"OK — Cubiertos: {cubiertos}/{total}, Ausentes: {len(ausencias)}, Tardanzas: {len(tardanzas)}")
